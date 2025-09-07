from __future__ import annotations

import importlib.util
from pathlib import Path
from unittest.mock import MagicMock

import pytest

from src.services.document_service import DocumentService


@pytest.fixture
def mocks():
    dense = MagicMock()
    lexical = MagicMock()
    dense.index_corpus.return_value = (
        ["1", "2"],
        {"status": "success", "count": 2},
    )
    lexical.index_documents.return_value = (
        ["1", "2"],
        {"status": "success", "count": 2},
    )
    dense.update_document.return_value = {"status": "success"}
    dense.delete_document.return_value = {"status": "success"}
    dense.validate_index.return_value = (True, {})
    lexical.update_document.return_value = {"status": "success"}
    lexical.delete_document.return_value = {"status": "success"}
    return dense, lexical


def create_pdf(path: Path) -> None:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="Hello PDF", ln=True)
    pdf.output(path)


def create_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Hello DOCX")
    doc.save(path)


def test_parse_various_formats(tmp_path: Path, mocks) -> None:
    dense, lexical = mocks
    service = DocumentService(dense, lexical)

    txt_file = tmp_path / "sample.txt"
    txt_file.write_text("hello txt")
    assert "hello txt" in service.parse_document(str(txt_file))

    html_file = tmp_path / "sample.html"
    html_file.write_text("<html><body><p>Hello <b>HTML</b></p></body></html>")
    assert "Hello HTML" in service.parse_document(str(html_file))

    md_file = tmp_path / "sample.md"
    md_file.write_text("# Title\n\nHello MD")
    parsed_md = service.parse_document(str(md_file))
    assert "Title" in parsed_md and "Hello MD" in parsed_md

    if importlib.util.find_spec("fpdf") and importlib.util.find_spec("pypdf"):
        pdf_file = tmp_path / "sample.pdf"
        create_pdf(pdf_file)
        assert "Hello PDF" in service.parse_document(str(pdf_file))

    if importlib.util.find_spec("docx"):
        docx_file = tmp_path / "sample.docx"
        create_docx(docx_file)
        assert "Hello DOCX" in service.parse_document(str(docx_file))


def test_chunk_and_ingest(tmp_path: Path, mocks) -> None:
    dense, lexical = mocks
    service = DocumentService(dense, lexical, chunk_size=3, overlap=1)
    file = tmp_path / "text.txt"
    file.write_text("one two three four five")
    result = service.ingest([str(file)])

    expected_chunks = ["one two three", "three four five"]
    assert dense.index_corpus.call_args[0][0] == expected_chunks
    assert lexical.index_documents.call_args[0][0] == expected_chunks
    assert result["dense"]["count"] == 2
    assert result["lexical"]["count"] == 2
    assert result["chunk_count"] == 2


def test_ingest_progress_callback(tmp_path: Path, mocks) -> None:
    dense, lexical = mocks
    service = DocumentService(dense, lexical, chunk_size=3, overlap=1)
    file = tmp_path / "text.txt"
    file.write_text("one two three four five")
    steps: list[str] = []

    def cb(pct: float, desc: str) -> None:
        steps.append(desc)

    service.ingest([str(file)], progress=cb)
    assert any("Parsing" in s for s in steps)
    assert any("Indexing dense" in s for s in steps)
    assert any("Indexing lexical" in s for s in steps)


def test_update_delete_and_audit(mocks) -> None:
    dense, lexical = mocks
    service = DocumentService(dense, lexical)

    service.update_document("1", "new content")
    service.delete_document("1")

    dense.update_document.assert_called_with("1", "new content", {})
    lexical.update_document.assert_called_with("1", "new content")
    dense.delete_document.assert_called_with("1")
    lexical.delete_document.assert_called_with("1")

    audit = service.audit_operations()
    assert [entry["action"] for entry in audit] == ["update", "delete"]


def test_bulk_operations_audit(mocks) -> None:
    dense, lexical = mocks
    service = DocumentService(dense, lexical)

    ops = [
        {"action": "update", "doc_id": "1", "content": "a"},
        {"action": "delete", "doc_id": "2"},
    ]
    service.bulk_operations(ops)

    dense.update_document.assert_called_with("1", "a", {})
    lexical.update_document.assert_called_with("1", "a")
    dense.delete_document.assert_called_with("2")
    lexical.delete_document.assert_called_with("2")
    audit = service.audit_operations()
    assert len(audit) == 2
